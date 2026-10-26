from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_THEME_COLOR
from pathlib import Path

OUT = Path(__file__).resolve().parents[1] / "submission" / "ReleaseProof_AI_Hackathon_Deck_v2.pptx"
OUT.parent.mkdir(exist_ok=True)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(22, 19, 31)
DARK = RGBColor(13, 11, 18)
INK = RGBColor(245, 241, 232)
MUTED = RGBColor(183, 176, 166)
LIME = RGBColor(183, 255, 60)
VIOLET = RGBColor(155, 123, 255)
CORAL = RGBColor(255, 128, 102)
TEAL = RGBColor(103, 232, 211)

def rect(slide, x, y, w, h, fill, radius=False, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    shape.line.fill.background()
    return shape

def text(slide, value, x, y, w, h, size, color=INK, bold=False, font="Aptos", align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]; p.alignment = align
    for i, line in enumerate(value.split("\n")):
        if i: p = tf.add_paragraph(); p.alignment = align
        run = p.add_run(); run.text = line
        run.font.name = font; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = color
    return box

def line(slide, x1, y1, x2, y2, color, width=1.5):
    sh = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh.line.color.rgb = color; sh.line.width = Pt(width)
    return sh

def base(section, n):
    s = prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb = NAVY
    rect(s, 0, 0, 13.333, .12, LIME)
    text(s, "RELEASEPROOF AI", .55, 7.08, 2.5, .16, 8, MUTED, True)
    text(s, section, 9.2, 7.08, 3.0, .16, 8, MUTED, False, align=PP_ALIGN.RIGHT)
    text(s, f"0{n}", 12.55, 7.05, .35, .18, 9, LIME, True, align=PP_ALIGN.RIGHT)
    return s

def pill(s, label, x, y, w, color):
    rect(s, x, y, w, .3, color, True)
    text(s, label, x+.1, y+.06, w-.2, .14, 8, DARK, True, align=PP_ALIGN.CENTER)

# 1
s=base("HACKATHON SUBMISSION",1)
text(s,"Every release should\nshow its work.",.7,1.1,7.2,1.4,30,INK,True)
text(s,"ReleaseProof AI",.72,3.15,5.5,.35,20,LIME,True)
text(s,"Evidence-driven release intelligence for Copado Agentia",.74,3.78,7,.3,14,MUTED)
rect(s,8.45,1.05,3.65,4.8,DARK,True)
text(s,"RELEASE\nDECISION",8.85,1.5,2.8,.5,13,MUTED,True)
text(s,"INSUFFICIENT",8.85,2.3,2.8,.3,18,CORAL,True)
line(s,8.85,2.9,11.72,2.9,MUTED,.7)
text(s,"Boundary case missing\nPermission check stale\nEvidence export ready",8.85,3.25,2.8,1.1,13,INK)
pill(s,"LOCAL DEMO",8.85,5.08,1.35,LIME)
text(s,"A transparent layer around Plan -> Build -> Test",.74,5.75,6.8,.25,12,TEAL)

# 2
s=base("THE GAP",2)
text(s,"Green tests can still miss the business rule",.7,.75,11,.5,24,INK,True)
text(s,"Teams can ship a passing build without proving the behavior that matters.",.72,1.48,10.5,.3,13,MUTED)
text(s,"WHAT TYPICAL RELEASE SIGNALS SHOW",.72,2.25,4.8,.2,9,MUTED,True)
text(s,"✓ happy path passes\n✓ build is green\n✓ deployment is available",.72,2.62,4.5,1.25,18,INK)
text(s,"WHAT THEY CAN MISS",6.9,2.25,4.8,.2,9,CORAL,True)
text(s,"× boundary equality\n× permission behavior\n× evidence invalidated by a change",6.9,2.62,5.4,1.25,18,INK)
line(s,.72,4.35,12.45,4.35,MUTED,.6)
text(s,"ReleaseProof makes the missing proof visible before approval.",.72,4.75,10,.45,22,LIME,True)
text(s,"The product does not replace Copado. It makes the release decision explainable.",.72,5.55,10.6,.3,13,MUTED)

# 3
s=base("USE CASE",3)
text(s,"US-247: lead routing at the equality boundary",.7,.75,11.5,.5,23,INK,True)
text(s,"Business contract",.72,1.55,2.3,.22,10,LIME,True)
text(s,"Score greater than 80 -> Senior Sales\nScore 80 or below -> Standard Sales\nOnly Sales Managers can override",.72,1.92,5.5,1.3,17,INK)
text(s,"Existing examples",7.05,1.55,2.5,.22,10,VIOLET,True)
text(s,"60 -> Standard Sales\n100 -> Senior Sales",7.05,1.92,4.5,.75,17,INK)
text(s,"Missing proof",.72,4.05,2.3,.22,10,CORAL,True)
text(s,"The suite omits score = 80 and does not verify the override permission.",.72,4.42,10,.45,20,CORAL,True)
text(s,"That is the exact gap ReleaseProof surfaces.",.72,5.33,8.8,.3,14,MUTED)

# 4
s=base("LIVE DEMO STORY",4)
text(s,"One change. One decision. A reason judges can see.",.7,.75,11.8,.5,23,INK,True)
xs=[.8,4.55,8.3]; titles=["1  ANALYZE","2  CHECK","3  RE-RUN"]; bodies=["Evidence incomplete\n\nThe contract needs a boundary check.","80 -> Senior Sales\n\nThe modeled behavior fails the equality case.","Evidence satisfied\n\nThe boundary fix closes the gap."]; cols=[VIOLET,CORAL,LIME]
for i in range(3):
    rect(s,xs[i],1.8,3.15,3.5,DARK,True); text(s,titles[i],xs[i]+.28,2.18,2.5,.22,10,cols[i],True); text(s,bodies[i],xs[i]+.28,2.72,2.55,1.5,17,INK)
    if i<2: text(s,"->",xs[i]+3.32,3.1,.35,.35,22,MUTED,True)
text(s,"The screen recording follows this sequence in LOCAL DEMO mode.",.8,5.95,10.8,.3,13,TEAL)

# 5
s=base("PRODUCT",5)
text(s,"The evidence ledger keeps the release honest",.7,.75,11.5,.5,23,INK,True)
for x,head,body in [(0.75,"INPUT","Business contract\nwork item\ntarget environment"),(4.42,"EVIDENCE","required checks\npass / fail / stale\nreason for confidence"),(8.2,"DECISION","satisfied\ninsufficient\nneeds new evidence")]:
    text(s,head,x,1.72,2,.2,9,MUTED,True); text(s,body,x,2.08,3.1,1.1,17,INK)
line(s,3.35,2.65,4.15,2.65,LIME,2); line(s,7.05,2.65,7.9,2.65,LIME,2)
rect(s,.75,4.42,10.8,1.15,DARK,True); text(s,"Exportable proof",1.05,4.72,2.1,.22,11,LIME,True); text(s,"Every decision carries its checks, evidence version, assumptions, and limitations.",3.2,4.7,7.7,.3,15,INK)

# 6
s=base("ARCHITECTURE",6)
text(s,"A local-first engine with a live Copado path",.7,.75,11.8,.5,23,INK,True)
for x,head,body,col in [(0.8,"RELEASEPROOF CORE","React / Next.js PWA\nFastAPI backend\nDeterministic evidence engine\nExportable JSON ledger",LIME),(5.02,"AGENTIA ADAPTER","CLI capability probe\nCICD work item context\nLocal quality gates\nCRT build polling",VIOLET),(9.25,"COPADO SERVICES","Pipelines and work items\nSalesforce environments\nCopado Robotic Testing\nAI and org context",TEAL)]:
    rect(s,x,1.8,3.25,3.85,DARK,True); text(s,head,x+.3,2.15,2.6,.22,10,col,True); text(s,body,x+.3,2.65,2.6,1.5,17,INK)
line(s,4.08,3.7,4.92,3.7,LIME,2); line(s,8.3,3.7,9.15,3.7,LIME,2)
text(s,"Local mode is fully reproducible. Live mode activates when the Copado organization is connected.",.8,6.12,11.2,.3,13,MUTED)

# 7
s=base("WHY IT MATTERS",7)
text(s,"The novelty is the decision layer",.7,.75,10.8,.5,24,INK,True)
text(s,"ReleaseProof adds the missing connection between intent and evidence.",.72,1.48,10.5,.3,13,MUTED)
items=[("1","Business-aware checks","Tests map back to the contract, not only to code coverage.",.78,2.3,LIME), ("2","Evidence has a lifecycle","A change can make yesterday's proof stale.",.78,3.65,VIOLET), ("3","Explainable release gates","The decision includes the reason and the next action.",6.8,2.3,CORAL), ("4","Local-first, integration-ready","Teams can prove behavior before credentials are available.",6.8,3.65,TEAL)]
for num,head,body,x,y,col in items:
    text(s,num,x,y,.4,.45,28,col,True); text(s,head,x+.67,y+.02,4.8,.26,16,INK,True); text(s,body,x+.67,y+.45,4.9,.35,12,MUTED)

# 8
s=base("AGENTIA HEADLESS FIT",8)
text(s,"ReleaseProof turns Agentia commands into release evidence",.7,.75,11.8,.5,23,INK,True)
text(s,"The CLI is the control surface. ReleaseProof is the decision layer around it.",.72,1.48,10.8,.3,13,MUTED)
steps=[("PLAN","agentia cicd work create","Create a work item from release intent",VIOLET),("BUILD","git commit -am \"feat: fix routing\"","Keep the change traceable in Git",TEAL),("TEST","agentia cicd work test --local","Run local quality gates before promotion",CORAL),("PROVE","releaseproof evidence evaluate","Map results to the business contract",LIME)]
for i,(head,cmd,desc,col) in enumerate(steps):
    x=.75+i*3.05
    rect(s,x,2.1,2.55,2.8,DARK,True); text(s,head,x+.22,2.4,2.1,.22,10,col,True); text(s,cmd,x+.22,2.85,2.1,.65,13,INK,True); text(s,desc,x+.22,3.78,2.1,.65,11,MUTED)
    if i<3: text(s,"->",x+2.65,3.2,.3,.3,20,MUTED,True)
rect(s,.75,5.45,11.05,.72,RGBColor(30,27,39),True); text(s,"Built in this project: Agentia CLI adapter, capability probes, local quality-gate path, CRT build polling, and a dry-run trace for the demo.",1.0,5.68,10.5,.22,12,INK)

# 9
s=base("BUSINESS IMPACT",9)
text(s,"The business case is fewer blind releases",.7,.75,11.5,.5,24,INK,True)
text(s,"ReleaseProof creates value before and after a deployment decision.",.72,1.48,10.5,.3,13,MUTED)
cols=[(.8,"BEFORE RELEASE",CORAL,"Less manual evidence hunting\nFewer missed boundary cases\nFaster go / no-go reviews"),(4.55,"DURING RELEASE",VIOLET,"One explainable decision\nStale evidence is visible\nAgentia checks stay traceable"),(8.3,"AFTER RELEASE",LIME,"Incident patterns become checks\nAudit trail is exportable\nLearning feeds the next release")]
for x,head,col,body in cols:
    rect(s,x,2.05,3.15,2.65,DARK,True); text(s,head,x+.25,2.35,2.6,.22,10,col,True); text(s,body,x+.25,2.88,2.6,1.15,16,INK)
text(s,"ROI model",.8,5.35,1.5,.22,11,TEAL,True)
text(s,"ROI = avoided escaped defects + reduced triage time + faster evidence collection - operating cost",2.15,5.31,9.7,.3,16,INK,True)
text(s,"Teams can plug in their own incident cost, release volume, and QA hours. No invented benchmark is required to measure the value.",.8,5.88,10.9,.3,12,MUTED)

# 10
s=base("TAKEAWAY",10)
text(s,"Ship with evidence,\nnot assumptions.",.75,1.0,7.2,1.3,30,INK,True)
text(s,"ReleaseProof AI",.78,2.85,4.2,.3,18,LIME,True)
text(s,"A practical release intelligence layer for Copado Agentia.",.8,3.4,7.3,.35,14,MUTED)
rect(s,8.4,1.35,3.5,3.5,LIME,True); text(s,"JUDGE\nTAKEAWAY",8.82,1.78,2.7,.45,12,DARK,True); text(s,"Did we prove\nthe contract?",8.82,2.65,2.7,.75,22,DARK,True)
text(s,"Demo: localhost:3000\nMode: LOCAL DEMO\nAgentia adapter: integrated and ready for org credentials",.8,5.2,6.6,.85,13,TEAL); text(s,"Thank you",8.82,5.6,2.5,.3,16,INK,True)

prs.save(OUT)
print(OUT)
